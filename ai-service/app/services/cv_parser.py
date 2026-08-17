from __future__ import annotations

import base64
import binascii
import re
from datetime import date
from io import BytesIO
from pathlib import Path

from app.schemas.cv import CVEducation, CVExperience, CVParseRequest, CVProfile, CVProject
from app.services.data_cleaning import (
    clean_text,
    extract_known_skills,
    extract_years_experience,
    normalize_location,
    normalize_skills,
)
from app.services.entity_extraction import normalize_newlines, split_text_sections


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?\d[\d\s().-]{7,}\d)(?!\d)")
DATE_RANGE_RE = re.compile(
    r"(?P<start>\d{4}(?:[-/]\d{1,2})?)\s*(?:to|-|–|—)\s*"
    r"(?P<end>present|current|now|\d{4}(?:[-/]\d{1,2})?)",
    re.IGNORECASE,
)
YEAR_RE = re.compile(r"\b(19\d{2}|20\d{2})\b")
DEGREE_PATTERNS = (
    ("Bachelor", r"\b(bachelor|b\.?s\.?|b\.?sc\.?|cu nhan|ky su)\b"),
    ("Master", r"\b(master|m\.?s\.?|m\.?sc\.?|thac si)\b"),
    ("PhD", r"\b(phd|doctor|tien si)\b"),
    ("Diploma", r"\b(diploma|cao dang)\b"),
)


class CVParseError(ValueError):
    """Raised when a CV cannot be decoded or parsed."""


def parse_cv_request(payload: CVParseRequest) -> CVProfile:
    if payload.text is not None:
        text = payload.text
    else:
        try:
            file_bytes = base64.b64decode(payload.file_base64 or "", validate=True)
        except (binascii.Error, ValueError) as exc:
            raise CVParseError("file_base64 is not valid base64") from exc
        text = extract_text_from_file_bytes(payload.filename or "", file_bytes)

    return parse_cv_text(
        text,
        include_raw_text=payload.include_raw_text,
        mask_pii=payload.mask_pii,
    )


def extract_text_from_path(path: str | Path) -> str:
    file_path = Path(path)
    return extract_text_from_file_bytes(file_path.name, file_path.read_bytes())


def extract_text_from_file_bytes(filename: str, file_bytes: bytes) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        return extract_text_from_pdf(file_bytes)
    if suffix == ".docx":
        return extract_text_from_docx(file_bytes)
    if suffix in {".txt", ".text"}:
        return file_bytes.decode("utf-8-sig")
    raise CVParseError("Unsupported CV file type. Use PDF, DOCX, or TXT.")


def extract_text_from_pdf(file_bytes: bytes) -> str:
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover - dependency is pinned in requirements
        raise CVParseError("PyMuPDF is required to parse PDF files") from exc

    try:
        with fitz.open(stream=file_bytes, filetype="pdf") as document:
            return "\n".join(page.get_text("text").strip() for page in document)
    except Exception as exc:  # noqa: BLE001 - library raises broad exceptions
        raise CVParseError("Could not read PDF text") from exc


def extract_text_from_docx(file_bytes: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover - dependency is pinned in requirements
        raise CVParseError("python-docx is required to parse DOCX files") from exc

    try:
        document = Document(BytesIO(file_bytes))
    except Exception as exc:  # noqa: BLE001 - library raises broad exceptions
        raise CVParseError("Could not read DOCX text") from exc

    parts: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def parse_cv_text(
    text: str,
    *,
    include_raw_text: bool = False,
    mask_pii: bool = True,
) -> CVProfile:
    normalized_text = normalize_cv_text(text)
    masked_text = mask_pii_text(normalized_text)
    source_text = masked_text if mask_pii else normalized_text
    sections = {section.name: section.text for section in split_text_sections(source_text)}
    header_lines = extract_header_lines(source_text)

    summary = clean_multiline(sections.get("summary"))
    skills_section = sections.get("skills")
    experiences = parse_experiences(sections.get("experience", ""))
    education = parse_education(sections.get("education", ""))
    projects = parse_projects(sections.get("projects", ""))
    certificates = parse_certificates(sections.get("certificates", ""))

    return CVProfile(
        candidate_name=extract_candidate_name(header_lines),
        title=extract_title(header_lines, experiences),
        summary=summary or None,
        email=None if mask_pii else extract_email(normalized_text),
        phone=None if mask_pii else extract_phone(normalized_text),
        location=extract_location(header_lines),
        skills=parse_profile_skills(skills_section, source_text),
        languages=parse_languages(sections.get("languages", "")),
        total_years_experience=extract_years_experience(source_text),
        experiences=experiences,
        education=education,
        projects=projects,
        certificates=certificates,
        raw_text=source_text if include_raw_text else None,
    )


def normalize_cv_text(text: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in normalize_newlines(text).splitlines()]
    compact_lines: list[str] = []
    blank_seen = False
    for line in lines:
        if not line:
            if not blank_seen and compact_lines:
                compact_lines.append("")
            blank_seen = True
            continue
        compact_lines.append(line)
        blank_seen = False
    return "\n".join(compact_lines).strip()


def mask_pii_text(text: str) -> str:
    masked = EMAIL_RE.sub("[email]", text)
    return PHONE_RE.sub("[phone]", masked)


def extract_header_lines(text: str) -> list[str]:
    sections = split_text_sections(text)
    header = next((section.text for section in sections if section.name == "header"), "")
    return [line.strip() for line in header.splitlines() if line.strip()]


def extract_candidate_name(header_lines: list[str]) -> str | None:
    for line in header_lines:
        if is_contact_or_link(line) or looks_like_location(line):
            continue
        return line
    return None


def extract_title(
    header_lines: list[str],
    experiences: list[CVExperience],
) -> str | None:
    content_lines = [
        line
        for line in header_lines
        if not is_contact_or_link(line) and not looks_like_location(line)
    ]
    if len(content_lines) >= 2:
        return content_lines[1]
    if experiences:
        return experiences[0].title
    return None


def extract_location(header_lines: list[str]) -> str | None:
    for line in header_lines:
        if looks_like_location(line):
            return normalize_location(line)
    return None


def looks_like_location(value: str) -> bool:
    normalized = normalize_location(value)
    return bool(
        normalized
        and (
            normalized != value
            or any(
                token in value.lower()
                for token in ("city", "remote", "hanoi", "ha noi", "ho chi minh", "da nang")
            )
        )
    )


def is_contact_or_link(value: str) -> bool:
    lowered = value.lower()
    return bool(
        EMAIL_RE.search(value)
        or PHONE_RE.search(value)
        or "://" in value
        or "[email]" in lowered
        or "[phone]" in lowered
    )


def parse_experiences(text: str) -> list[CVExperience]:
    lines = [line for line in text.splitlines() if line.strip()]
    experiences: list[CVExperience] = []
    index = 0

    while index < len(lines):
        header = lines[index].strip()
        date_match = DATE_RANGE_RE.search(header)
        if date_match is None:
            index += 1
            continue

        description_lines: list[str] = []
        next_index = index + 1
        while next_index < len(lines) and DATE_RANGE_RE.search(lines[next_index]) is None:
            description_lines.append(lines[next_index].strip())
            next_index += 1

        title, company = parse_experience_header(header, date_match)
        description = clean_multiline("\n".join(description_lines))
        skill_text = " ".join([header, description or ""])
        experiences.append(
            CVExperience(
                title=title,
                company=company,
                start_date=parse_date_token(date_match.group("start")),
                end_date=parse_date_token(date_match.group("end")),
                description=description or None,
                skills=normalize_skills(extract_known_skills(skill_text)),
            )
        )
        index = next_index

    return experiences


def parse_profile_skills(skills_section: str | None, fallback_text: str) -> list[str]:
    if skills_section:
        declared_skills = normalize_skills(skills_section)
        if declared_skills:
            return declared_skills
    return normalize_skills(extract_known_skills(fallback_text))


def parse_experience_header(
    header: str,
    date_match: re.Match[str],
) -> tuple[str, str | None]:
    prefix = header[: date_match.start()].strip(" ,-")
    parts = [part.strip() for part in prefix.split(",") if part.strip()]
    if not parts:
        return "Experience", None
    title = parts[0]
    company = parts[1] if len(parts) > 1 else None
    return title, company


def parse_education(text: str) -> list[CVEducation]:
    education: list[CVEducation] = []
    for line in [line.strip() for line in text.splitlines() if line.strip()]:
        degree, major = parse_degree_and_major(line)
        years = [int(match.group(1)) for match in YEAR_RE.finditer(line)]
        end_year = years[-1] if years else None
        start_year = years[0] if len(years) > 1 else infer_start_year(degree, end_year)
        school = parse_school(line, degree, major, end_year)
        if not school and not degree:
            continue
        education.append(
            CVEducation(
                school=school or "Unknown",
                degree=degree,
                major=major,
                start_year=start_year,
                end_year=end_year,
            )
        )
    return education


def parse_degree_and_major(line: str) -> tuple[str | None, str | None]:
    degree = None
    for canonical, pattern in DEGREE_PATTERNS:
        if re.search(pattern, line, flags=re.IGNORECASE):
            degree = canonical
            break

    major = None
    major_match = re.search(
        r"\b(?:of|in)\s+([^,\n]+?)(?:,|\bat\b|\bfrom\b|\d{4}|$)",
        line,
        flags=re.IGNORECASE,
    )
    if major_match:
        major = clean_text(major_match.group(1))
    return degree, major


def parse_school(
    line: str,
    degree: str | None,
    major: str | None,
    end_year: int | None,
) -> str | None:
    cleaned = line
    if end_year is not None:
        cleaned = cleaned.replace(str(end_year), "")
    if degree:
        cleaned = re.sub(degree, "", cleaned, flags=re.IGNORECASE)
    if major:
        cleaned = cleaned.replace(f"of {major}", "").replace(f"in {major}", "")
    parts = [part.strip(" ,.-") for part in cleaned.split(",") if part.strip(" ,.-")]
    return parts[-1] if parts else None


def infer_start_year(degree: str | None, end_year: int | None) -> int | None:
    if end_year is None:
        return None
    if degree == "Bachelor":
        return end_year - 4
    if degree == "Master":
        return end_year - 2
    return None


def parse_projects(text: str) -> list[CVProject]:
    projects: list[CVProject] = []
    for line in [line.strip(" -") for line in text.splitlines() if line.strip(" -")]:
        if ":" in line:
            name, description = [part.strip() for part in line.split(":", 1)]
        elif " - " in line:
            name, description = [part.strip() for part in line.split(" - ", 1)]
        else:
            name, description = line, None
        projects.append(
            CVProject(
                name=name,
                description=description,
                skills=normalize_skills(extract_known_skills(description or name)),
            )
        )
    return projects


def parse_certificates(text: str) -> list[str]:
    return [line.strip(" -") for line in text.splitlines() if line.strip(" -")]


def parse_languages(text: str) -> list[str]:
    known = {
        "english": "English",
        "vietnamese": "Vietnamese",
        "tieng anh": "English",
        "tieng viet": "Vietnamese",
    }
    found: list[str] = []
    folded = text.lower()
    for alias, canonical in known.items():
        if alias in folded and canonical not in found:
            found.append(canonical)
    return found


def parse_date_token(value: str) -> date | None:
    token = value.strip().lower()
    if token in {"present", "current", "now"}:
        return None
    match = re.match(r"(?P<year>\d{4})(?:[-/](?P<month>\d{1,2}))?", token)
    if not match:
        return None
    year = int(match.group("year"))
    month = int(match.group("month") or 1)
    return date(year, month, 1)


def clean_multiline(text: str | None) -> str:
    return clean_text(text or "")


def extract_email(text: str) -> str | None:
    match = EMAIL_RE.search(text)
    return match.group(0) if match else None


def extract_phone(text: str) -> str | None:
    match = PHONE_RE.search(text)
    return match.group(0) if match else None
